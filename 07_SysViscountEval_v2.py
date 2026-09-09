"""Compare manual counts with detections; update SQLite and the Excel report."""
import ast
import atexit
from copy import copy
import os
from pathlib import Path
import sys
import tempfile

import numpy as np
import pandas as pd

from utils.funcs_vis_count_noseller_time import short_name, visitors_counting
from utils.funcs_initializer_camconfig_getcamframe import load_camconfig, save_camconfig
import utils.db as db
from utils.funcs_TxtUI_request_app_description import cleanup_mei_folders, get_app_name

atexit.register(cleanup_mei_folders)

DESCRIPTION = (
    'Оценка точности подсчета посетителей; требуется инициализированная db/cv.db.\n'
    'При наличии db/1_real_viscount.xlsx данные берутся из его листов\n'
    'по короткому имени точки (chm для chm1); иначе — из SQLite.\n'
    'Результат: SQLite и db/1_Sys_viscount_eval.xlsx.\n'
    'Неполные ручные дни и дни без детекций пропускаются с пояснением в статусе.\n'
    'Формат: имя_основной_камеры: (mean_threshold, window_next), например chm1: (1, 12).\n'
    'mean_threshold >= 0, window_next >= 1; оба значения целые.\n'
)


def algorithm_params(value):
    try:
        pair = ast.literal_eval(value) if isinstance(value, str) else value
        if (not isinstance(pair, (tuple, list)) or len(pair) != 2
                or any(type(v) is not int for v in pair) or pair[0] < 0 or pair[1] < 1):
            raise ValueError
        return tuple(pair)
    except (SyntaxError, ValueError, TypeError):
        raise ValueError('Ожидается пара целых (mean >= 0, window >= 1)') from None


def is_main_camera(name):
    return bool(name) and (not name[-1].isdigit() or name[-1] == '1')


def create_txt_params(text_path, camconfig, description):
    with open(text_path, 'w', encoding='utf-8') as f:
        for line in description.strip().splitlines():
            print(f'# {line}', file=f)
        for cam in camconfig:
            if is_main_camera(cam['cam_name']):
                print(f"{cam['cam_name']}: {cam['vis_count_alg']}", file=f)


def create_txt_program_status(text_note_path, msg):
    Path(text_note_path).write_text(msg + '\n', encoding='utf-8')


def read_txt_params(text_note_path):
    params = {}
    for number, line in enumerate(Path(text_note_path).read_text(encoding='utf-8-sig').splitlines(), 1):
        line = line.split('#', 1)[0].strip()
        if not line:
            continue
        try:
            name, value = line.split(':', 1)
            name = name.strip()
            if not name or name in params:
                raise ValueError('Пустое или повторное имя камеры')
            params[name] = algorithm_params(value.strip())
        except ValueError as error:
            raise ValueError(f'Строка {number}: {error}') from error
    if not params:
        raise ValueError('В файле параметров не указаны камеры')
    return params


def normalize_manual(frame, hours, notify):
    """Reject incomplete/ambiguous days instead of treating missing counts as zero."""
    frame = frame.copy()
    frame.columns = [str(c).strip() for c in frame.columns]
    if frame.empty:
        return pd.DataFrame(columns=['date'] + hours + ['sum', '*'])
    if 'date' not in frame or frame.columns.duplicated().any():
        raise ValueError('Ручные данные: отсутствует date или повторяются столбцы')
    frame = frame.dropna(how='all')
    dates = pd.to_datetime(frame['date'], errors='coerce', format='mixed')
    numeric_dates = frame['date'].map(lambda v: isinstance(v, (int, float, np.number)))
    try:
        frame['date'] = dates.dt.normalize()
        if dates.dt.tz is not None:
            raise ValueError('Даты должны быть без часового пояса')
    except AttributeError:
        raise ValueError('Некорректный формат дат в ручных данных') from None
    values = frame.reindex(columns=hours).apply(pd.to_numeric, errors='coerce')
    valid = (np.isfinite(values) & values.ge(0) & values.eq(values.round())).all(axis=1)
    valid &= frame['date'].notna() & ~numeric_dates & ~frame['date'].duplicated(keep=False)
    rejected = int((~valid).sum())
    if rejected:
        notify(f'Пропущено ручных строк: {rejected} (дата, дубликат, пропуск часа или неверное число)')
    result = values.loc[valid].astype('int64')
    result.insert(0, 'date', frame.loc[valid, 'date'])
    result['sum'] = result[hours].sum(axis=1)
    result['*'] = frame.loc[valid, '*'] if '*' in frame else None
    return result.sort_values('date').reset_index(drop=True)


def evaluation(cam_name, params, camconfig, cwd_path, manual=None, notify=print):
    """Recalculate eligible manual days; return report rows, or None when skipped."""
    config = next((c for c in camconfig if c['cam_name'] == cam_name), None)
    if config is None or not is_main_camera(cam_name):
        raise ValueError(f'Неизвестная или дополнительная камера: {cam_name}')
    mean, window = algorithm_params(params[cam_name])
    start, end = ast.literal_eval(config['work_hours'])
    if not (0 <= start < end <= 24):
        raise ValueError(f'{cam_name}: неверные рабочие часы')
    hours = [str(h) for h in range(start, end)]
    shop = short_name(cam_name)
    if manual is None:
        manual = db.read_real_viscount(shop, cwd_path, preserve_missing=True)
    real = normalize_manual(manual, hours, notify)
    if real.empty:
        notify('Нет полных ручных дней; прежние результаты сохранены')
        return None

    shapes = db.read_shapes(cam_name, cwd_path)
    shapes = shapes.sort_values('origin_file_name').reset_index(drop=True)
    shapes['cam_name'] = cam_name
    # Read and group once, rather than loading the entire history per day.
    day_groups = {day: group for day, group in shapes.groupby(shapes['origin_file_name'].str[:6])}
    pairs = []
    for _, row in real.iterrows():
        day = row['date'].strftime('%y%m%d')
        day_shapes = day_groups.get(day)
        if day_shapes is None or day_shapes.empty:
            notify(f'{row["date"]:%Y-%m-%d}: нет детекций; день пропущен')
            continue
        auto = visitors_counting(cam_name, day_shapes, day, mean, window, cwd_path=cwd_path)
        auto_row = {h: int(auto.iloc[0].get(h, 0)) for h in hours}
        auto_sum = sum(auto_row.values())
        real_sum = int(row['sum'])
        error = real_sum - auto_sum
        mape = round(abs(error) / real_sum, 2) if real_sum else (0.0 if auto_sum == 0 else None)
        if mape is None:
            notify(f'{row["date"]:%Y-%m-%d}: ручной итог 0, автоматический {auto_sum}; MAPE не определен')
        common = {'date': row['date'], '*': row['*'], 'err': error,
                  'mape': 'n/a' if mape is None else str(mape).replace('.', ',')}
        pairs.extend([
            {**row.to_dict(), **common, 's': 'real'},
            {**auto_row, **common, 'sum': auto_sum, 's': 'auto'},
        ])
    if not pairs:
        notify('Нет дней для оценки; прежние результаты и параметры сохранены')
        return None
    result = pd.DataFrame(pairs)[['date'] + hours + ['sum', '*', 's', 'err', 'mape']]
    db.write_evstat(shop, result, cwd_path, mode='replace')

    visitors = db.read_visitors(shop, cwd_path)
    if not visitors.empty:
        visitors['date'] = pd.to_datetime(visitors['date']).dt.normalize()
        replacements = real[real['date'].isin(visitors['date'])].copy()
        replacements['s'] = 'real'
        if not replacements.empty:
            visitors = visitors.set_index('date')
            replacements = replacements.set_index('date')
            for hour in hours:
                if hour not in visitors:
                    visitors[hour] = 0
            visitors.update(replacements)
            visitors['sum'] = visitors[[c for c in visitors if str(c).isdigit()]].sum(axis=1)
            db.write_visitors(shop, visitors.reset_index(), cwd_path, mode='replace')
    # Import validated rows only, retaining previously stored other dates.
    db.write_real_viscount(shop, real, cwd_path, mode='append')
    return result


def export_excel(path, reports):
    """Update table cells in place; preserve sheets, styles and summary formulas."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path = Path(path)
    workbook = load_workbook(path) if path.exists() else Workbook()
    if not path.exists():
        workbook.remove(workbook.active)
    temporary = None
    try:
        for shop, frame in reports.items():
            title = f'{shop}_evstat'
            sheet = workbook[title] if title in workbook.sheetnames else workbook.create_sheet(title)
            columns = list(frame.columns)
            if sheet.cell(1, 1).value is not None:
                headers = [sheet.cell(1, col).value for col in range(1, len(columns) + 1)]
                if [str(v) for v in headers] != columns:
                    raise ValueError(f'{title}: столбцы шаблона не совпадают с рабочими часами камеры')
            for col, name in enumerate(columns, 1):
                sheet.cell(1, col, name)
            for row in sheet.iter_rows(min_row=2, max_row=max(sheet.max_row, len(frame) + 1), max_col=len(columns)):
                for cell in row:
                    cell.value = None
            for i, values in enumerate(frame.itertuples(index=False, name=None), 2):
                for j, value in enumerate(values, 1):
                    cell = sheet.cell(i, j)
                    if i > 2:
                        cell._style = copy(sheet.cell(2, j)._style)
                    if columns[j - 1] == 'mape':
                        value = None if value == 'n/a' else float(str(value).replace(',', '.'))
                        cell.number_format = '0.00%'
                    elif pd.isna(value):
                        value = None
                    elif isinstance(value, pd.Timestamp):
                        value = value.to_pydatetime()
                        cell.number_format = 'yyyy-mm-dd'
                    cell.value = value
                    if columns[j - 1] == '*' and isinstance(value, str):
                        cell.data_type = 's'
            ref = f'A1:{get_column_letter(len(columns))}{len(frame) + 1}'
            table = sheet.tables.get(title)
            if table is None:
                table = Table(displayName=title, ref=ref)
                table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
                sheet.add_table(table)
            table.ref = ref
            if table.autoFilter is not None:
                table.autoFilter.ref = ref
        workbook.calculation.fullCalcOnLoad = True
        path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as f:
            temporary = Path(f.name)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main(cwd_path=None):
    root = Path(cwd_path or os.getcwd())
    app_name = get_app_name()
    request = root / f'{app_name}_request_app_description.txt'
    status = root / f'{app_name}_program_status.txt'
    messages = []
    failed = False
    try:
        create_txt_program_status(status, 'Ждите...')
        camconfig = load_camconfig(root)
        if not request.exists():
            create_txt_params(request, camconfig, DESCRIPTION)
            create_txt_program_status(status, 'Файл параметров создан. Проверьте его и запустите программу повторно.')
            return 0
        params = read_txt_params(request)
        reports = {}
        manual_path = root / 'db' / '1_real_viscount.xlsx'
        manual_sheets = pd.read_excel(manual_path, sheet_name=None) if manual_path.exists() else None
        for cam_name in params:
            def notify(message):
                text = f'{cam_name}: {message}'
                messages.append(text)
                print(text)
            try:
                if not is_main_camera(cam_name) or not any(c['cam_name'] == cam_name for c in camconfig):
                    raise ValueError(f'Неизвестная или дополнительная камера: {cam_name}')
                if manual_sheets is not None and short_name(cam_name) not in manual_sheets:
                    notify('Нет листа ручных данных; камера пропущена')
                    continue
                manual = None if manual_sheets is None else manual_sheets[short_name(cam_name)]
                result = evaluation(cam_name, params, camconfig, root, manual=manual, notify=notify)
                if result is not None:
                    reports[short_name(cam_name)] = result
                    for config in camconfig:
                        if short_name(config['cam_name']) == short_name(cam_name):
                            config['vis_count_alg'] = params[cam_name]
            except Exception as error:
                failed = True
                notify(f'Ошибка: {error}')
        if reports:
            save_camconfig(camconfig, root)
            export_excel(root / 'db' / '1_Sys_viscount_eval.xlsx', reports)
            messages.append(f'Обновлен Excel-отчет; точек: {len(reports)}')
        heading = 'Завершено с ошибками' if failed else ('Готово' if reports else 'Нет данных для оценки')
        create_txt_program_status(status, heading + '\n' + '\n'.join(messages))
        return 1 if failed else 0
    except Exception as error:
        message = f'Ошибка: {error}'
        if isinstance(error, db.DatabaseIdentityError):
            message += '\nНужна база Retail db/cv.db; старую CSV-копию сначала импортируйте migrate_csv_to_sqlite.py.'
        if isinstance(error, PermissionError):
            message += '\nЗакройте Excel и проверьте права доступа, затем повторите запуск. SQLite мог уже обновиться.'
        create_txt_program_status(status, message + '\n' + '\n'.join(messages))
        print(message)
        return 1


if __name__ == '__main__':
    sys.exit(main())
