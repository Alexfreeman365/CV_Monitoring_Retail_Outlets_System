"""Regression tests use isolated SQLite databases and real counting/Excel code."""
import importlib.util
from contextlib import closing
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from utils import db

spec = importlib.util.spec_from_file_location('sysviscount', Path(__file__).resolve().parents[1] / '07_SysViscountEval_v2.py')
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


class EvaluationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        db.init_db(self.root)
        self.config = [dict(cam_name='chm1', shape_zone='(0, 10, 0, 10)',
                            face_zone='(0, 10, 0, 10)', frame='(0, 10, 0, 10)',
                            work_hours='(10, 19)', vis_count_alg='(1, 12)')]
        db.save_camconfig(self.config, self.root)
        self.hours = [str(h) for h in range(10, 19)]
        self.params = {'chm1': (1, 12)}
        self.messages = []

    def manual(self, count=2, date='2024-09-25'):
        return pd.DataFrame([{'date': date, **{h: count for h in self.hours}, '*': 'note'}])

    def shapes(self, zone=1):
        files = ['24092510000010.jpg', '24092510010010.jpg', '24092510010010.jpg',
                 '24092510010010.jpg', '24092510020010.jpg']
        data = pd.DataFrame([dict(origin_file_name=f, uid8=f'240925{i}', shape_location='[0, 10, 0, 10]',
                                 shape_zone_coords='(0, 10, 0, 10)', shape_zone=zone,
                                 face_zone_coords='(0, 10, 0, 10)', face_zone=0) for i, f in enumerate(files)])
        db.write_shapes('chm1', data, self.root)

    def evaluate(self, manual=None):
        return app.evaluation('chm1', self.params, self.config, self.root,
                              manual=manual, notify=self.messages.append)

    def test_empty_manual_is_skipped(self):
        self.assertIsNone(self.evaluate())
        self.assertTrue(self.messages)

    def test_missing_shapes_preserve_existing_results(self):
        self.shapes()
        self.evaluate(self.manual())
        before = db.read_evstat('chm', self.root)
        with closing(sqlite3.connect(db.db_path(self.root))) as conn, conn:
            conn.execute('DELETE FROM shapes_locs')
        self.params['chm1'] = (2, 12)
        self.assertIsNone(self.evaluate(self.manual()))
        pd.testing.assert_frame_equal(before, db.read_evstat('chm', self.root))

    def test_missing_hour_and_invalid_values_are_skipped(self):
        self.shapes()
        for value in [None, -1, 1.5, float('inf'), 'oops']:
            with self.subTest(value=value):
                data = self.manual().astype(object)
                data.loc[0, '18'] = value
                self.assertIsNone(self.evaluate(data))
        self.assertIsNone(self.evaluate(self.manual().drop(columns='18')))

    def test_invalid_duplicate_dates_and_numeric_dates(self):
        self.shapes()
        for date in ['bad', 1234]:
            self.assertIsNone(self.evaluate(self.manual(date=date)))
        self.assertIsNone(self.evaluate(pd.concat([self.manual(), self.manual()], ignore_index=True)))

    def test_sqlite_missing_hour_in_one_day_is_not_zero(self):
        self.shapes()
        data = pd.concat([self.manual(), self.manual(date='2024-09-26')], ignore_index=True)
        db.write_real_viscount('chm', data, self.root)
        with closing(sqlite3.connect(db.db_path(self.root))) as conn, conn:
            conn.execute("DELETE FROM real_viscount WHERE date='2024-09-25' AND hour=18")
        self.assertIsNone(self.evaluate())
        self.assertTrue(any('Пропущено ручных строк: 1' in m for m in self.messages))

    def test_manual_replaces_existing_visitors_and_preserves_other_dates(self):
        self.shapes()
        data = pd.concat([self.manual(99), self.manual(99, '2024-09-26')], ignore_index=True)
        data['s'] = 'auto'
        db.write_visitors('chm', data, self.root)
        result = self.evaluate(self.manual())
        self.assertEqual(len(result), 2)
        saved = db.read_visitors('chm', self.root).set_index('date')
        self.assertEqual(saved.loc['2024-09-25', 'sum'], 18)
        self.assertEqual(saved.loc['2024-09-25', 's'], 'real')
        self.assertEqual(saved.loc['2024-09-26', 'sum'], 891)
        self.assertEqual(saved.loc['2024-09-26', 's'], 'auto')

    def test_zero_manual_positive_auto_roundtrip(self):
        self.shapes()
        result = self.evaluate(self.manual(0))
        self.assertGreater(result.loc[result.s == 'auto', 'sum'].iloc[0], 0)
        self.assertEqual(set(db.read_evstat('chm', self.root)['mape']), {'n/a'})
        path = self.root / 'out.xlsx'
        app.export_excel(path, {'chm': result})
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertIsNone(book['chm_evstat']['O2'].value)

    def test_zero_both(self):
        self.shapes(zone=0)
        self.evaluate(self.manual(0))
        self.assertEqual(set(db.read_evstat('chm', self.root)['mape']), {'0,0'})

    def test_repeat_and_modified_manual_recalculate(self):
        self.shapes()
        self.evaluate(self.manual())
        before = db.read_evstat('chm', self.root)
        self.evaluate(self.manual())
        pd.testing.assert_frame_equal(before, db.read_evstat('chm', self.root))
        self.evaluate(self.manual(3))
        current = db.read_evstat('chm', self.root)
        self.assertEqual(current.loc[current.s == 'real', 'sum'].iloc[0], 27)
        self.assertEqual(len(current), 2)
        self.params['chm1'] = (2, 3)
        self.assertEqual(len(self.evaluate(self.manual())), 2)

    def test_excel_preserves_summary_and_other_sheets(self):
        self.shapes()
        result = self.evaluate(self.manual())
        path = self.root / 'out.xlsx'
        app.export_excel(path, {'chm': result})
        book = load_workbook(path)
        book['chm_evstat']['Q1'] = '=AVERAGE(O:O)'
        book.create_sheet('unrelated')['A1'] = 'keep'
        book.save(path)
        book.close()
        app.export_excel(path, {'chm': result})
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book['chm_evstat']['Q1'].value, '=AVERAGE(O:O)')
        self.assertEqual(book['unrelated']['A1'].value, 'keep')
        self.assertEqual(book['chm_evstat'].tables['chm_evstat'].ref, 'A1:O3')

    def test_locked_excel_preserves_previous_file(self):
        self.shapes()
        result = self.evaluate(self.manual())
        path = self.root / 'out.xlsx'
        app.export_excel(path, {'chm': result})
        original = path.read_bytes()
        with patch.object(app.os, 'replace', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                app.export_excel(path, {'chm': result})
        self.assertEqual(original, path.read_bytes())
        self.assertEqual(list(self.root.glob('*.xlsx')), [path])

    def test_parameter_validation(self):
        path = self.root / 'params.txt'
        path.write_text('\ufeff# описание\n chm1 : (1,12) # comment\n', encoding='utf-8')
        self.assertEqual(app.read_txt_params(path), self.params)
        for content in ['chm1: (1, 0)', 'chm1: (-1, 2)', 'chm1: (1.5, 2)',
                        'chm1: (1, 2)\nchm1: (2, 3)', 'broken', '# empty']:
            path.write_text(content, encoding='utf-8')
            with self.subTest(content=content), self.assertRaises(ValueError):
                app.read_txt_params(path)

    def test_main_reads_excel_edits_each_run(self):
        self.shapes()
        with patch.object(app, 'get_app_name', return_value='eval'):
            self.assertEqual(app.main(self.root), 0)
            for count in [2, 3]:
                self.manual(count).to_excel(self.root / 'db/1_real_viscount.xlsx', sheet_name='chm', index=False)
                self.assertEqual(app.main(self.root), 0)
                frame = db.read_evstat('chm', self.root)
                self.assertEqual(frame.loc[frame.s == 'real', 'sum'].iloc[0], 9 * count)
            self.assertTrue((self.root / 'db/1_Sys_viscount_eval.xlsx').exists())

    def test_main_reports_missing_database_without_traceback(self):
        with patch.object(app, 'get_app_name', return_value='eval'):
            empty = self.root / 'empty'
            empty.mkdir()
            self.assertEqual(app.main(empty), 1)
            self.assertIn('migrate_csv_to_sqlite', (empty / 'eval_program_status.txt').read_text(encoding='utf-8'))

    def test_main_continues_after_unknown_camera(self):
        self.shapes()
        db.write_real_viscount('chm', self.manual(), self.root)
        (self.root / 'eval_request_app_description.txt').write_text('unknown: (1, 12)\nchm1: (1, 12)', encoding='utf-8')
        with patch.object(app, 'get_app_name', return_value='eval'):
            self.assertEqual(app.main(self.root), 1)
        self.assertEqual(len(db.read_evstat('chm', self.root)), 2)
        self.assertTrue((self.root / 'db/1_Sys_viscount_eval.xlsx').exists())
        self.assertIn('Завершено с ошибками', (self.root / 'eval_program_status.txt').read_text(encoding='utf-8'))

    def test_missing_excel_sheet_does_not_use_cached_manual_data(self):
        self.shapes()
        db.write_real_viscount('chm', self.manual(), self.root)
        self.manual().to_excel(self.root / 'db/1_real_viscount.xlsx', sheet_name='chm_arc', index=False)
        with patch.object(app, 'get_app_name', return_value='eval'):
            app.main(self.root)
            self.assertEqual(app.main(self.root), 0)
        self.assertTrue(db.read_evstat('chm', self.root).empty)

    def test_main_failed_export_can_be_retried(self):
        self.shapes()
        db.write_real_viscount('chm', self.manual(), self.root)
        with patch.object(app, 'get_app_name', return_value='eval'):
            app.main(self.root)
            with patch.object(app.os, 'replace', side_effect=PermissionError('locked')):
                self.assertEqual(app.main(self.root), 1)
            self.assertNotIn('Ждите...', (self.root / 'eval_program_status.txt').read_text(encoding='utf-8'))
            self.assertEqual(app.main(self.root), 0)
        self.assertTrue((self.root / 'db/1_Sys_viscount_eval.xlsx').exists())


if __name__ == '__main__':
    unittest.main()
